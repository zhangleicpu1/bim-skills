# -*- coding: utf-8 -*-
"""报告生成：Excel 明细表 / Word 分析报告 / Markdown 摘要。

对外接口：
    build_markdown(analysis, out_dir)  -> path
    build_excel(analysis, out_dir)     -> path
    build_docx(analysis, out_dir)      -> path
"""

from __future__ import annotations

import os
from datetime import datetime

SEVERITY_FILL = {"严重": "F8CBAD", "高": "FFE699", "中": "DDEBF7", "低": "E2EFDA"}
SEVERITY_HEX = {"P0": "#C0392B", "P1": "#E67E22", "P2": "#F1C40F", "P3": "#27AE60"}
CN_FONT = "微软雅黑"


def _safe(name: str) -> str:
    return "".join("_" if c in '\\/:*?"<>|' else c for c in name).strip() or "报告"


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _stem(analysis: dict) -> str:
    src = os.path.splitext(os.path.basename(analysis["source"]))[0]
    return _safe(src)


# ==========================================================================
# 图表（matplotlib，缺失或字体不可用时静默跳过）
# ==========================================================================
def _setup_font():
    try:
        from matplotlib import font_manager, rcParams
    except Exception:
        return None
    candidates = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC",
                  "PingFang SC", "Hiragino Sans GB", "WenQuanYi Zen Hei", "Arial Unicode MS"]
    available = {f.name for f in font_manager.fontManager.ttflist}
    for c in candidates:
        if c in available:
            rcParams["font.sans-serif"] = [c]
            rcParams["axes.unicode_minus"] = False
            return c
    return None


def make_charts(analysis: dict, out_dir: str, prefix: str) -> list[str]:
    """生成 2 张汇总图，返回文件路径列表。失败时返回空列表。"""
    font = _setup_font()
    if not font:
        return []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return []

    s = analysis["summary"]
    paths: list[str] = []

    # ---- 图 1：严重程度 + 碰撞类型 ----
    fig, axes = plt.subplots(1, 2, figsize=(10, 3.6), dpi=160)
    rows = s["by_severity"]
    labels = [f"{r['label']}\n({r['key']})" for r in rows]
    counts = [r["count"] for r in rows]
    colors = [SEVERITY_HEX.get(r["key"], "#8E9BA8") for r in rows]
    ax = axes[0]
    bars = ax.bar(labels, counts, color=colors, width=0.6)
    ax.set_title("严重程度分布", fontsize=11)
    ax.set_ylabel("碰撞点数", fontsize=9)
    ax.tick_params(labelsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    for b, c in zip(bars, counts):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + max(counts + [1]) * 0.02,
                str(c), ha="center", fontsize=8)

    ax = axes[1]
    trows = [r for r in s["by_clash_type"] if r["count"]]
    ax.pie([r["count"] for r in trows],
           labels=[f"{r['key']} {r['count']}" for r in trows],
           autopct="%1.1f%%", textprops={"fontsize": 8},
           colors=["#E74C3C", "#F39C12", "#3498DB", "#95A5A6"][:len(trows)])
    ax.set_title("碰撞类型占比", fontsize=11)
    plt.tight_layout()
    p1 = os.path.join(out_dir, f"{prefix}_chart_severity.png")
    fig.savefig(p1, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths.append(p1)

    # ---- 图 2：专业组合 + 楼层 ----
    pair_rows = [r for r in s["by_discipline_pair"] if r["key"] != "未知"][:10]
    level_rows = [r for r in s["by_level"] if r["key"] != "未标注"][:10]
    if pair_rows or level_rows:
        fig, axes = plt.subplots(1, 2, figsize=(10, 3.6), dpi=160)

        ax = axes[0]
        if pair_rows:
            names = [r["key"] for r in pair_rows][::-1]
            vals = [r["count"] for r in pair_rows][::-1]
            ax.barh(names, vals, color="#4A90C4", height=0.6)
            ax.set_title("专业组合 TOP10", fontsize=11)
            ax.tick_params(labelsize=8)
            for y, v in enumerate(vals):
                ax.text(v + max(vals) * 0.015, y, str(v), va="center", fontsize=8)
        else:
            ax.axis("off")
        ax.spines[["top", "right"]].set_visible(False)

        ax = axes[1]
        if level_rows:
            names = [r["key"] for r in level_rows]
            vals = [r["count"] for r in level_rows]
            ax.bar(names, vals, color="#7FB77E", width=0.6)
            ax.set_title("楼层分布 TOP10", fontsize=11)
            ax.tick_params(labelsize=8, rotation=30)
            for x, v in enumerate(vals):
                ax.text(x, v + max(vals) * 0.02, str(v), ha="center", fontsize=8)
        else:
            ax.axis("off")
        ax.spines[["top", "right"]].set_visible(False)

        plt.tight_layout()
        p2 = os.path.join(out_dir, f"{prefix}_chart_distribution.png")
        fig.savefig(p2, bbox_inches="tight", facecolor="white")
        plt.close(fig)
        paths.append(p2)

    return paths


# ==========================================================================
# Markdown 摘要
# ==========================================================================
def build_markdown(analysis: dict, out_dir: str) -> str:
    s = analysis["summary"]
    records = analysis["records"]
    lines = [
        f"# BIM 碰撞检测分析摘要",
        "",
        f"- **数据来源**：`{os.path.basename(analysis['source'])}`（{analysis['format'].upper()} 格式）",
        f"- **生成时间**：{_stamp()}",
        f"- **判级规则**：{analysis['rules_version']}",
        f"- **碰撞点总数**：{s['total']}",
        f"- **需整改（严重+高）**：{s['need_action']}",
        "",
        "## 一、严重程度分布",
        "",
        "| 等级 | 数量 | 占比 |",
        "| --- | ---: | ---: |",
    ]
    for r in s["by_severity"]:
        lines.append(f"| {r['label']}（{r['key']}） | {r['count']} | {r['ratio']:.1%} |")

    lines += ["", "## 二、碰撞类型分布", "", "| 类型 | 数量 | 占比 |", "| --- | ---: | ---: |"]
    for r in s["by_clash_type"]:
        lines.append(f"| {r['key']} | {r['count']} | {r['ratio']:.1%} |")

    lines += ["", "## 三、专业组合 TOP", "", "| 专业组合 | 数量 |", "| --- | ---: |"]
    for r in s["by_discipline_pair"][:10]:
        lines.append(f"| {r['key']} | {r['count']} |")

    lines += ["", "## 四、楼层分布", "", "| 楼层 | 数量 |", "| --- | ---: |"]
    for r in s["by_level"][:12]:
        lines.append(f"| {r['key']} | {r['count']} |")

    lines += ["", "## 五、重点碰撞清单（严重 + 高）", "",
              "| 编号 | 严重度 | 分值 | 类型 | 专业组合 | 楼层 | 构件 A | 构件 B | 距离 |",
              "| --- | --- | ---: | --- | --- | --- | --- | --- | ---: |"]
    rank = {"P0": 0, "P1": 1}
    keyed = [r for r in records if r["severity_grade"] in rank]
    keyed.sort(key=lambda r: (rank[r["severity_grade"]], -r["severity_score"]))
    for r in keyed[:30]:
        lines.append("| {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
            r["name"], r["severity"], r["severity_score"], r["clash_type"],
            r["discipline_pair"], r["level"],
            (r["item1"]["name"] or "—")[:36], (r["item2"]["name"] or "—")[:36],
            r["distance_str"]))

    lines += ["", "## 六、高频冲突构件 TOP15", "", "| 构件 | 专业 | 碰撞次数 | 最严重等级 |", "| --- | --- | ---: | --- |"]
    for r in s["top_items"]:
        lines.append(f"| {(r['name'] or '—')[:48]} | {r['discipline']} | {r['count']} | {r['worst']} |")

    path = os.path.join(out_dir, f"{_stem(analysis)}_碰撞分析摘要.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return path


# ==========================================================================
# Excel
# ==========================================================================
def build_excel(analysis: dict, out_dir: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    s = analysis["summary"]
    records = analysis["records"]

    head_font = Font(name=CN_FONT, bold=True, size=10, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    body_font = Font(name=CN_FONT, size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(ws, ncols, row=1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

    def autosize(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # ---------- Sheet 1：分析概览 ----------
    ws = wb.active
    ws.title = "分析概览"
    ws["A1"] = "BIM 碰撞检测分析概览"
    ws["A1"].font = Font(name=CN_FONT, bold=True, size=14)
    ws.merge_cells("A1:D1")
    info = [
        ("数据来源", os.path.basename(analysis["source"])),
        ("源文件格式", analysis["format"].upper()),
        ("报告生成时间", _stamp()),
        ("判级规则版本", analysis["rules_version"]),
        ("碰撞点总数", s["total"]),
        ("待处理（未关闭）", s["open_count"]),
        ("需整改（严重+高）", s["need_action"]),
    ]
    r = 3
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(name=CN_FONT, bold=True, size=10)
        ws.cell(row=r, column=2, value=v).font = body_font
        r += 1

    r += 1
    blocks = [
        ("严重程度分布", ["等级", "数量", "占比"],
         [[f"{x['label']}（{x['key']}）", x["count"], f"{x['ratio']:.1%}"] for x in s["by_severity"]]),
        ("碰撞类型分布", ["类型", "数量", "占比"],
         [[x["key"], x["count"], f"{x['ratio']:.1%}"] for x in s["by_clash_type"]]),
        ("专业组合分布", ["专业组合", "数量", ""],
         [[x["key"], x["count"], ""] for x in s["by_discipline_pair"]]),
        ("楼层分布", ["楼层", "数量", ""],
         [[x["key"], x["count"], ""] for x in s["by_level"]]),
    ]
    for title, header, body in blocks:
        ws.cell(row=r, column=1, value=title).font = Font(name=CN_FONT, bold=True, size=11)
        r += 1
        for c, h in enumerate(header, 1):
            ws.cell(row=r, column=c, value=h)
        style_header(ws, len(header), row=r)
        r += 1
        for row in body:
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = body_font
                cell.border = border
            r += 1
        r += 1
    autosize(ws, [40, 14, 12, 10])

    # ---------- Sheet 2：碰撞明细 ----------
    ws2 = wb.create_sheet("碰撞明细")
    cols = ["编号", "碰撞测试项", "状态", "是否待处理", "严重等级", "严重度分值", "碰撞类型",
            "专业组合", "构件类别组合", "楼层", "轴网", "距离", "碰撞点坐标",
            "构件A名称", "构件A专业", "构件A类别", "构件A类型", "构件A路径",
            "构件B名称", "构件B专业", "构件B类别", "构件B类型", "构件B路径",
            "判级依据", "整改建议", "责任人", "备注"]
    ws2.append(cols)
    style_header(ws2, len(cols))

    sorted_records = sorted(records, key=lambda r: (
        {"P0": 0, "P1": 1, "P2": 2, "P3": 3}[r["severity_grade"]], -r["severity_score"]))
    for i, rec in enumerate(sorted_records, 1):
        ws2.append([
            i, rec["_clash_test"], rec["status"], "是" if rec["is_open"] else "否",
            rec["severity"], rec["severity_score"], rec["clash_type"],
            rec["discipline_pair"], rec["category_pair"], rec["level"], rec.get("grid", ""),
            rec["distance_str"], rec["clash_point_str"],
            rec["item1"]["name"], rec["item1"]["discipline"], rec["item1"]["category"],
            rec["item1"]["type"], rec["item1"]["path"],
            rec["item2"]["name"], rec["item2"]["discipline"], rec["item2"]["category"],
            rec["item2"]["type"], rec["item2"]["path"],
            rec["severity_reason"], rec["suggestion"], "", rec.get("comment", ""),
        ])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=len(cols)):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=False)
        sev = row[4].value
        if sev in SEVERITY_FILL:
            row[4].fill = PatternFill("solid", fgColor=SEVERITY_FILL[sev])
            row[4].alignment = Alignment(horizontal="center", vertical="center")
        row[5].alignment = Alignment(horizontal="center", vertical="center")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws2.max_row}"
    autosize(ws2, [6, 20, 9, 10, 9, 10, 10, 20, 18, 10, 10, 12, 22, 38, 16, 12, 22, 40,
                   38, 16, 12, 22, 40, 60, 60, 10, 16])

    # ---------- Sheet 3：重点整改清单 ----------
    ws3 = wb.create_sheet("重点整改清单")
    head3 = ["序号", "严重等级", "分值", "碰撞测试项", "专业组合", "楼层", "碰撞类型",
             "构件A", "构件B", "距离", "整改建议", "责任专业", "责任人", "计划完成", "实际完成", "状态"]
    ws3.append(head3)
    style_header(ws3, len(head3))
    keyed = [r for r in sorted_records if r["severity_grade"] in ("P0", "P1")]
    for i, rec in enumerate(keyed, 1):
        ws3.append([
            i, rec["severity"], rec["severity_score"], rec["_clash_test"], rec["discipline_pair"],
            rec["level"], rec["clash_type"], rec["item1"]["name"], rec["item2"]["name"],
            rec["distance_str"], rec["suggestion"], rec["discipline_pair"], "", "", "", "待整改",
        ])
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=len(head3)):
        for cell in row:
            cell.font = body_font
            cell.border = border
        sev = row[1].value
        if sev in SEVERITY_FILL:
            row[1].fill = PatternFill("solid", fgColor=SEVERITY_FILL[sev])
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(head3))}{ws3.max_row}"
    autosize(ws3, [6, 9, 8, 20, 20, 10, 10, 36, 36, 12, 55, 20, 10, 12, 12, 10])

    # ---------- Sheet 4：统计汇总 ----------
    ws4 = wb.create_sheet("统计汇总")
    r = 1
    for title, header, body in [
        ("按严重程度", ["等级", "数量", "占比"], [[f"{x['label']}（{x['key']}）", x["count"], f"{x['ratio']:.1%}"] for x in s["by_severity"]]),
        ("按碰撞类型", ["类型", "数量", "占比"], [[x["key"], x["count"], f"{x['ratio']:.1%}"] for x in s["by_clash_type"]]),
        ("按专业组合", ["专业组合", "数量"], [[x["key"], x["count"]] for x in s["by_discipline_pair"]]),
        ("按构件类别组合", ["类别组合", "数量"], [[x["key"], x["count"]] for x in s["by_category_pair"]]),
        ("按专业（单侧）", ["专业", "涉及构件数"], [[x["key"], x["count"]] for x in s["by_discipline"]]),
        ("按楼层", ["楼层", "数量"], [[x["key"], x["count"]] for x in s["by_level"]]),
        ("按碰撞测试项", ["测试项", "数量"], [[x["key"], x["count"]] for x in s["by_test"]]),
        ("按状态", ["状态", "数量"], [[x["key"], x["count"]] for x in s["by_status"]]),
        ("高频冲突构件 TOP15", ["构件", "专业", "次数", "最严重等级"],
         [[x["name"], x["discipline"], x["count"], x["worst"]] for x in s["top_items"]]),
    ]:
        ws4.cell(row=r, column=1, value=title).font = Font(name=CN_FONT, bold=True, size=11)
        r += 1
        for c, h in enumerate(header, 1):
            ws4.cell(row=r, column=c, value=h)
        style_header(ws4, len(header), row=r)
        r += 1
        for row in body:
            for c, v in enumerate(row, 1):
                cell = ws4.cell(row=r, column=c, value=v)
                cell.font = body_font
                cell.border = border
            r += 1
        r += 1
    autosize(ws4, [46, 14, 12, 14])

    # ---------- Sheet 5：判级规则说明 ----------
    ws5 = wb.create_sheet("判级规则说明")
    ws5["A1"] = "严重程度判级规则"
    ws5["A1"].font = Font(name=CN_FONT, bold=True, size=12)
    lines = analysis["rules_explain"]
    for i, text in enumerate(lines, 3):
        c = ws5.cell(row=i, column=1, value=text)
        c.font = body_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws5, [120])
    for i in range(3, 3 + len(lines)):
        ws5.row_dimensions[i].height = 30

    path = os.path.join(out_dir, f"{_stem(analysis)}_碰撞明细.xlsx")
    wb.save(path)
    return path


# ==========================================================================
# Word
# ==========================================================================
def build_docx(analysis: dict, out_dir: str, charts: list[str] | None = None) -> str:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    s = analysis["summary"]
    records = analysis["records"]
    project = analysis.get("project") or {}
    title_text = project.get("project") or analysis.get("title") or _stem(analysis)

    def set_font(run, size=10.5, bold=False, color=None, name=CN_FONT):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), name)

    def para(text="", size=10.5, bold=False, color=None, align=None, space_after=6, indent=False):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.74)
        if text:
            set_font(p.add_run(text), size=size, bold=bold, color=color)
        return p

    def heading(text, level=1):
        sizes = {1: 15, 2: 12.5, 3: 11.5}
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        p.paragraph_format.space_after = Pt(6)
        set_font(p.add_run(text), size=sizes.get(level, 11.5), bold=True,
                 color="1F4E79" if level == 1 else "2E5C8A")
        return p

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def set_cell(cell, text, size=9, bold=False, fill=None, align=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        if align:
            p.alignment = align
        set_font(p.add_run("" if text is None else str(text)), size=size, bold=bold)
        if fill:
            shade(cell, fill)

    def make_table(header, rows, widths=None, size=9):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            set_cell(t.rows[0].cells[i], h, size=size, bold=True, fill="DCE6F1",
                     align=WD_ALIGN_PARAGRAPH.CENTER)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row[:len(header)]):
                set_cell(cells[i], v, size=size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    if i < len(row.cells):
                        row.cells[i].width = Cm(w)
        return t

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    # ===== 封面标题 =====
    para("BIM 碰撞检测分析报告", size=22, bold=True, color="1F4E79",
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(title_text, size=12, color="595959", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    make_table(
        ["项目", "内容", "项目", "内容"],
        [
            ["数据来源", os.path.basename(analysis["source"]),
             "源文件格式", analysis["format"].upper()],
            ["碰撞点总数", s["total"], "待处理", s["open_count"]],
            ["需整改（严重+高）", s["need_action"], "涉及碰撞测试项",
             len([x for x in s["by_test"] if x["key"] != "未分组"]) or len(s["by_test"])],
            ["报告生成时间", _stamp(), "判级规则版本", analysis["rules_version"]],
        ],
        widths=[3.0, 5.4, 3.4, 4.2], size=9.5)

    # ===== 一、检测概况 =====
    heading("一、检测概况", 1)
    para(f"本报告基于 Navisworks 导出的碰撞检测结果（{os.path.basename(analysis['source'])}）"
         f"进行自动整理与分级。原始数据共包含 {s['total']} 个碰撞点，覆盖 "
         f"{len(s['by_test'])} 个碰撞测试项、{len(s['by_discipline_pair'])} 种专业组合、"
         f"{len(s['by_level'])} 个楼层。按本报告判级规则，其中严重（P0）{s['severity_counts'].get('P0', 0)} 处、"
         f"高（P1）{s['severity_counts'].get('P1', 0)} 处，合计 {s['need_action']} 处需优先整改，"
         f"占比 {s['need_action'] / s['total']:.1%}。" if s["total"] else "原始数据中未解析到有效碰撞记录。",
         indent=True)

    make_table(
        ["严重程度", "数量", "占比", "处理要求"],
        [[r["label"] + f"（{r['key']}）", r["count"], f"{r['ratio']:.1%}",
          {"P0": "立即处理，影响结构安全或主干系统，须在下一版模型前闭合",
           "P1": "优先处理，影响净高、检修空间或点位合规性",
           "P2": "计划处理，纳入常规碰撞协调会跟踪",
           "P3": "低优先，可结合深化设计一并解决"}.get(r["key"], "")]
         for r in s["by_severity"]],
        widths=[3.0, 2.0, 2.0, 9.0])

    # ===== 二、碰撞类型分类汇总 =====
    heading("二、碰撞类型分类汇总", 1)
    para("碰撞类型依据检测方式与 Navisworks 距离符号自动判定：负距离（存在穿透）为硬碰撞，"
         "正距离（存在净距但小于设定容差）为间隙碰撞，构件完全重合为重复项。", indent=True)
    make_table(["碰撞类型", "数量", "占比", "典型特征"],
               [[r["key"], r["count"], f"{r['ratio']:.1%}",
                 {"硬碰撞": "构件实体相互穿越，必须调整几何位置",
                  "间隙碰撞": "未直接穿透但净距不足，需复核检修/保温/安全净距",
                  "重复项": "构件重复建模，通常删除冗余模型即可",
                  "未知": "报告未提供判定信息"}.get(r["key"], "")]
                for r in s["by_clash_type"]],
               widths=[3.0, 2.0, 2.0, 9.0])

    # ===== 三、严重程度分级汇总 =====
    heading("三、严重程度分级汇总", 1)
    para("结合碰撞类型、构件专业优先级、穿透深度/净距、是否位于关键区域、是否跨专业等因子加权评分，"
         "并按阈值划分为四个等级。详细规则见文末附录。", indent=True)
    if charts:
        for cp in charts:
            doc.add_picture(cp, width=Cm(16.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ===== 四、分布特征 =====
    heading("四、碰撞分布特征", 1)
    heading("4.1 专业组合分布", 2)
    para("下表按专业组合统计碰撞数量，用于识别需要重点协调的专业界面。", indent=True)
    make_table(["专业组合", "数量", "占比", "主要矛盾"],
               [[r["key"], r["count"], f"{r['count'] / s['total']:.1%}" if s["total"] else "-",
                 analysis["pair_advice"].get(r["key"], "")[:60]]
                for r in s["by_discipline_pair"][:15]],
               widths=[4.0, 1.8, 1.8, 8.4])

    if len(s["by_level"]) > 1:
        heading("4.2 楼层分布", 2)
        make_table(["楼层", "数量", "占比"],
                   [[r["key"], r["count"], f"{r['count'] / s['total']:.1%}" if s["total"] else "-"]
                    for r in s["by_level"][:20]],
                   widths=[4.0, 2.0, 2.0])

    heading("4.3 构件类别组合 TOP10", 2)
    make_table(["类别组合", "数量"],
               [[r["key"], r["count"]] for r in s["by_category_pair"][:10]],
               widths=[6.0, 2.0])

    # ===== 五、重点碰撞清单 =====
    heading("五、重点碰撞清单（严重 + 高）", 1)
    rank = {"P0": 0, "P1": 1}
    keyed = sorted([r for r in records if r["severity_grade"] in rank],
                   key=lambda r: (rank[r["severity_grade"]], -r["severity_score"]))
    if keyed:
        para(f"共 {len(keyed)} 处碰撞被评定为严重或高等级，按分值降序排列如下（最多列出 50 条，"
             f"完整清单见配套 Excel 明细表）。", indent=True)
        rows = []
        for i, r in enumerate(keyed[:50], 1):
            rows.append([i, r["severity"], r["clash_type"], r["discipline_pair"], r["level"],
                         (r["item1"]["name"] or "—")[:26], (r["item2"]["name"] or "—")[:26],
                         r["distance_str"]])
        make_table(["#", "等级", "类型", "专业组合", "楼层", "构件 A", "构件 B", "距离"],
                   rows, widths=[0.9, 1.3, 1.7, 3.0, 1.5, 3.6, 3.6, 1.8], size=8)
    else:
        para("本次检测未发现严重或高等级碰撞。", indent=True)

    # ===== 六、高频冲突构件 =====
    heading("六、高频冲突构件 TOP15", 1)
    para("下列构件反复出现碰撞，通常提示其定位、系统标高或专业间接口存在系统性问题，"
         "建议整系统复核而非逐个处理。", indent=True)
    make_table(["构件名称", "所属专业", "碰撞次数", "最严重等级"],
               [[(r["name"] or "—")[:52], r["discipline"], r["count"], r["worst"]]
                for r in s["top_items"]],
               widths=[8.4, 3.0, 2.2, 2.4])

    # ===== 七、整改建议 =====
    heading("七、整改建议", 1)
    advice_groups: dict[str, dict] = {}
    for r in keyed:
        g = advice_groups.setdefault(r["discipline_pair"], {"count": 0, "p0": 0, "advice": r["suggestion"]})
        g["count"] += 1
        if r["severity_grade"] == "P0":
            g["p0"] += 1
    if not advice_groups:
        para("无需整改建议。", indent=True)
    for pair, g in sorted(advice_groups.items(), key=lambda kv: -kv[1]["count"]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"▍{pair}（{g['count']} 处，其中严重 {g['p0']} 处）"),
                 size=11, bold=True, color="2E5C8A")
        para(g["advice"] or analysis["default_remediation"], indent=True)

    para(f"整改闭环要求：以上问题应明确责任专业与责任人，设定计划完成时间；"
         f"模型更新后重新执行碰撞检测，确认对应碰撞点状态变为「已解决」后关闭。"
         f"建议将本报告配套的《碰撞明细》与《重点整改清单》作为碰撞协调会的过程文件。",
         indent=True)

    # ===== 八、附录 =====
    heading("八、附录：判级规则与统计口径", 1)
    for line in analysis["rules_explain"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run("• " + line), size=9)

    para(f"本报告由 BIM 碰撞检测报告整理工具自动生成，统计口径以源报告解析结果为准；"
         f"如源报告中缺少「碰撞类型 / 楼层」等信息，工具会依据距离符号与构件路径自动推断，"
         f"推断结果仅供参考，最终以设计协调结论为准。", size=9, color="595959")

    path = os.path.join(out_dir, f"{_stem(analysis)}_碰撞分析报告.docx")
    doc.save(path)
    return path
